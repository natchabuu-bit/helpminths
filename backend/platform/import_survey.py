"""Import the KPD three-row-header template into public aggregate statistics.

Run locally; never upload the original workbook to the public repository. Excel
formulas are read from saved values, never executed. Import fails before writing
if the layout, core counts, duplicate identities or source totals are ambiguous.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TRANSFORM_VERSION = 'kpd-workbook-aggregate-v2'
TYPE_COLUMNS = (
    ('hookworm', 'ปากขอ', 9), ('ascaris', 'ไส้เดือน', 11),
    ('trichuris', 'แส้ม้า', 13), ('liver_fluke', 'ใบไม้ตับ', 15),
    ('pinworm', 'เข็มหมุด', 17), ('tapeworm', 'ตืด', 19),
    ('medium_intestinal', 'ลำไส้ขนาดกลาง', 21), ('strongyloides', 'สตรองจิลอย', 23),
)
CORE_COLUMNS = {'enrolled': 4, 'tested': 5, 'infected': 7, 'coinfection': 25}


def _text(value) -> str:
    return unicodedata.normalize('NFC', str(value or '')).strip()


def _compact(value) -> str:
    return re.sub(r'\s+', '', _text(value))


def _count(value, cell: str) -> int | None:
    if value is None or value == '':
        return None
    if type(value) not in (int, float) or not math.isfinite(value) or value < 0 or value != int(value):
        raise ValueError(f'จำนวนใน {cell} ต้องเป็นจำนวนเต็มไม่ติดลบและมีค่าที่บันทึกไว้')
    return int(value)


def _sum_known(values):
    known = [value for value in values if value is not None]
    return sum(known) if known else None


def _positivity_band(infected: int | None, tested: int | None) -> str | None:
    """Descriptive bands of reporting units, never clinical risk categories."""
    if tested is None or tested <= 0 or infected is None:
        return None
    if infected == 0:
        return 'none'
    # Compare integers at the boundaries to avoid rounded percentage assignment.
    if infected * 100 < tested * 5:
        return 'under5'
    if infected * 100 < tested * 10:
        return '5to10'
    if infected * 100 < tested * 20:
        return '10to20'
    return '20plus'


def import_workbook(path: str | Path, *, worksheet: str, year: int, imported_at: str | None = None) -> dict:
    source = Path(path)
    if source.suffix.lower() != '.xlsx' or source.stat().st_size > 10 * 1024 * 1024:
        raise ValueError('รองรับไฟล์ .xlsx ขนาดไม่เกิน 10 MB')
    if not 2400 <= year <= 2743:
        raise ValueError('ระบุปีชุดข้อมูลเป็น พ.ศ.')
    if str(year) not in source.stem:
        raise ValueError('ปีชุดข้อมูลต้องตรงกับชื่อไฟล์ เพื่อไม่สร้างปีสำรวจที่ต้นฉบับไม่ได้ระบุ')
    with ZipFile(source) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('เนื้อหาไฟล์ที่คลายแล้วใหญ่เกินขอบเขตนำเข้า')
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb = load_workbook(source, data_only=True, read_only=True)
    formula_wb = load_workbook(source, data_only=False, read_only=True)
    try:
        if worksheet not in wb.sheetnames:
            raise ValueError('ไม่พบแท็บที่ระบุ')
        sheet, formulas = wb[worksheet], formula_wb[worksheet]
        if sheet.max_row > 10000 or sheet.max_column > 64:
            raise ValueError('ตารางใหญ่เกินขอบเขตแม่แบบที่รองรับ')
        rows = list(sheet.iter_rows(values_only=True))
        formula_rows = list(formulas.iter_rows(values_only=True))
        header = None
        for index, row in enumerate(rows[:30]):
            school = next((column for column, value in enumerate(row) if _compact(value) == 'โรงเรียน'), None)
            if school is not None and school + 27 < len(row) and _compact(row[school + 3]) == 'จังหวัด':
                header = (index, school)
                break
        if header is None:
            raise ValueError('ไม่พบหัวตารางสามแถวของข้อมูลสถานศึกษา')
        head, start = header
        if head + 2 >= len(rows):
            raise ValueError('หัวตารางต้องมีครบสามแถว')
        if _compact(rows[head][start + 1]) != 'ตำบล' or _compact(rows[head][start + 2]) != 'อำเภอ' or _compact(rows[head][start + 27]) != 'สังกัด':
            raise ValueError('ตำแหน่งคอลัมน์พื้นที่หรือสังกัดไม่ตรงกับแม่แบบ')
        if 'ทั้งหมด' not in _compact(rows[head][start + 4]) or 'ส่งตรวจ' not in _compact(rows[head][start + 5]) or 'ตรวจพบพยาธิ' not in _compact(rows[head][start + 7]):
            raise ValueError('หัวคอลัมน์จำนวนหลักไม่ตรงกับแม่แบบ')
        for _, label, offset in TYPE_COLUMNS:
            if _compact(rows[head + 1][start + offset]) != _compact(label) or 'จำนวน' not in _compact(rows[head + 2][start + offset]):
                raise ValueError('หัวคอลัมน์ชนิดพยาธิไม่ตรงกับแม่แบบ')
        if 'มากกว่า1ชนิด' not in _compact(rows[head + 1][start + 25]):
            raise ValueError('ไม่พบคอลัมน์การติดพยาธิมากกว่า 1 ชนิด')
        has_sequence = start > 0 and _compact(rows[head][start - 1]) == 'ที่'
        records, issues, identities = [], [], set()
        summary_start = None
        formula_error_count = 0

        def read_count(idx, offset):
            col = start + offset
            value = rows[idx][col]
            formula = formula_rows[idx][col]
            ref = f'{get_column_letter(col + 1)}{idx + 1}'
            if isinstance(formula, str) and formula.startswith('=') and value is None:
                raise ValueError(f'{ref} เป็นสูตรจำนวนที่ไม่มีค่าบันทึก ให้เปิดคำนวณและบันทึกด้วย Excel ก่อนนำเข้า')
            return _count(value, ref)

        for idx in range(head + 3, len(rows)):
            row = rows[idx]
            row_number = idx + 1
            if _compact(row[start + 3]) in {'รวม', 'รวมทั้งหมด', 'รวมทั้งสิ้น'}:
                summary_start = idx
                break
            name = _text(row[start])
            if not name:
                if any(value is not None for value in row):
                    raise ValueError(f'แถว {row_number} ไม่ใช่รายละเอียดหรือแถวรวมที่รองรับ')
                continue
            if has_sequence and row[start - 1] != len(records) + 1:
                raise ValueError(f'ลำดับหน่วยรายงานขาดหรือซ้ำที่แถว {row_number}')
            dims = dict(province=_text(row[start + 3]), district=_text(row[start + 2]), affiliation=_text(row[start + 27]))
            if not all(dims.values()):
                raise ValueError(f'พื้นที่หรือสังกัดว่างที่แถว {row_number}')
            identity = (name, _text(row[start + 1]), dims['district'], dims['province'], dims['affiliation'])
            if identity in identities:
                raise ValueError(f'พบหน่วยรายงานซ้ำที่แถว {row_number} ต้องตรวจต้นทางก่อนนำเข้า')
            identities.add(identity)
            core = {key: read_count(idx, offset) for key, offset in CORE_COLUMNS.items()}
            if all(value is None for value in core.values()):
                status = 'pending'
            elif any(value is None for value in core.values()):
                raise ValueError(f'จำนวนหลักรายงานไม่ครบที่แถว {row_number} ต้องตรวจต้นทางก่อนนำเข้า')
            else:
                status = 'tested' if core['tested'] > 0 else 'zero_tested'
                if not (core['coinfection'] <= core['infected'] <= core['tested'] <= core['enrolled']):
                    raise ValueError(f'จำนวนหลักขัดกันที่แถว {row_number} ต้องตรวจต้นทางก่อนนำเข้า')
            positivity_band = _positivity_band(core['infected'], core['tested'])
            row_issues = []
            if status == 'pending':
                row_issues.append({'code': 'pending_counts', 'severity': 'info', 'title': 'หน่วยรายงานรอข้อมูล', 'message': 'ยังไม่มีจำนวนนักเรียนและผลตรวจ ไม่แทนช่องว่างเป็นศูนย์', 'count': 1, 'cell_refs': [f'{get_column_letter(start + 5)}{row_number}', f'{get_column_letter(start + 6)}{row_number}', f'{get_column_letter(start + 8)}{row_number}']})
            parasites = {}
            for key, _, offset in TYPE_COLUMNS:
                ref = f'{get_column_letter(start + offset + 1)}{row_number}'
                value = read_count(idx, offset)
                if status == 'pending' and value is not None:
                    raise ValueError(f'มีจำนวนแยกชนิดแต่ไม่มีจำนวนหลักที่ {ref}')
                conflict = value is not None and core['infected'] is not None and value > core['infected']
                missing = int(value is None and status != 'pending')
                parasites[key] = {'reported': value, 'consistent': 0 if conflict else value, 'conflicting': value if conflict else 0,
                                  'missing': missing, 'tested_missing': missing if status == 'tested' else 0}
                if conflict:
                    row_issues.append({'code': 'source_count_conflict', 'severity': 'warning', 'parasite': key, 'title': 'จำนวนแยกชนิดขัดกับยอดรวม',
                                       'message': f'{ref} รายงาน {value} คน แต่ยอดส่งตรวจ {core["tested"]} และยอดตรวจพบ {core["infected"]} คน คงค่ารายงานไว้โดยแยกรายการที่ขัดกัน',
                                       'count': 1, 'reported': value, 'consistent': 0, 'cell_refs': [ref]})
                if missing:
                    row_issues.append({'code': 'missing_parasite_count', 'severity': 'warning', 'parasite': key, 'title': 'ช่องจำนวนแยกชนิดยังว่าง',
                                       'message': f'{ref} ไม่ระบุจำนวน ไม่ตีความว่าเป็นศูนย์ ผลรวมชนิดนี้อ้างเฉพาะช่องที่มีตัวเลข', 'count': 1, 'cell_refs': [ref]})
            percentage_offsets = [6, 8, *range(10, 27, 2)]
            error_refs = [f'{get_column_letter(start + col + 1)}{row_number}' for col in percentage_offsets if isinstance(row[start + col], str) and row[start + col].startswith('#')]
            if error_refs:
                formula_error_count += len(error_refs)
                row_issues.append({'code': 'formula_error', 'severity': 'info', 'title': 'สูตรร้อยละต้นทางไม่มีตัวหาร',
                                   'message': 'คำนวณร้อยละใหม่จากจำนวนคน กรณีตัวหารศูนย์หรือยังไม่มีข้อมูลจะแสดงว่าไม่สามารถคำนวณได้', 'count': len(error_refs), 'cell_refs': error_refs})
            for issue in row_issues:
                issues.append({'id': f'row-{row_number}-{issue["code"]}-{issue.get("parasite", "all")}', **dims,
                               'status': status, 'positivity_band': positivity_band, **issue})
            records.append({**dims, 'status': status, 'positivity_band': positivity_band,
                            **core, 'parasites': parasites, '_row': row_number})
        if not records or summary_start is None:
            raise ValueError('ไม่พบรายละเอียดหรือแถวรวมปิดตารางที่ใช้ตรวจยอด')
        check_offsets = {**CORE_COLUMNS, **{key: col for key, _, col in TYPE_COLUMNS}}

        def totals(selected):
            return {key: _sum_known([record[key] if key in CORE_COLUMNS else record['parasites'][key]['reported'] for record in selected]) for key in check_offsets}

        total_rows, province_rows = [], []
        global_totals = totals(records)
        provinces = {record['province'] for record in records}
        for idx in range(summary_start, len(rows)):
            row = rows[idx]
            province = _text(row[start + 3])
            if _text(row[start]) or not (province in provinces or _compact(province) in {'รวม', 'รวมทั้งหมด', 'รวมทั้งสิ้น'}):
                continue
            if all(row[start + offset] is None for offset in CORE_COLUMNS.values()):
                continue
            selected = records if _compact(province) in {'รวม', 'รวมทั้งหมด', 'รวมทั้งสิ้น'} else [record for record in records if record['province'] == province]
            expected = global_totals if selected is records else totals(selected)
            for key, offset in check_offsets.items():
                cell = f'{get_column_letter(start + offset + 1)}{idx + 1}'
                actual = read_count(idx, offset)
                if actual != expected[key]:
                    raise ValueError(f'ยอดสรุป {cell} ไม่ตรงกับผลรวมรายละเอียด หยุดนำเข้าเพื่อไม่เผยแพร่ยอดผิด')
            (total_rows if selected is records else province_rows).append(idx + 1)
        if not total_rows:
            raise ValueError('ไม่มีแถวรวมที่ตรวจสอบสำเร็จ')
        grouped = defaultdict(list)
        for record in records:
            grouped[tuple(record[key] for key in ('province', 'district', 'affiliation', 'status', 'positivity_band'))].append(record)
        cubes = []
        for dims, selected in sorted(grouped.items()):
            cubes.append({**dict(zip(('province', 'district', 'affiliation', 'status', 'positivity_band'), dims)),
                          'school_count': len(selected), 'positive_school_count': sum(bool(record['infected']) for record in selected),
                          **{key: _sum_known([record[key] for record in selected]) for key in CORE_COLUMNS},
                          'parasites': {key: {field: _sum_known([record['parasites'][key][field] for record in selected]) for field in ('reported', 'consistent', 'conflicting', 'missing', 'tested_missing')} for key, _, _ in TYPE_COLUMNS}})
        # Version the transformation as well as the workbook, so corrections are auditable.
        version_basis = json.dumps([digest, TRANSFORM_VERSION, worksheet, year, source.name], ensure_ascii=False)
        version = hashlib.sha256(version_basis.encode()).hexdigest()[:16]
        return {
            'schema_version': 1,
            'dataset': {'id': f'kpd-{year}', 'version': version, 'transform_version': TRANSFORM_VERSION,
                        'file_name': source.name, 'file_sha256': digest, 'worksheet': worksheet, 'year': year,
                        'year_basis': 'filename', 'band_partitioned': True,
                        'imported_at': imported_at or datetime.now(timezone.utc).isoformat(timespec='seconds')},
            'cubes': cubes, 'issues': issues,
            'audit': {'source_detail_rows': len(records), 'detail_row_range': [records[0]['_row'], records[-1]['_row']],
                      'excluded_summary_rows': total_rows, 'excluded_province_summary_rows': province_rows,
                      'summary_reconciled': True, 'formula_error_count': formula_error_count,
                      'formula_cells': sum(isinstance(cell, str) and cell.startswith('=') for row in formula_rows for cell in row),
                      'missing_count_cells': sum(record['parasites'][key]['missing'] for record in records for key, _, _ in TYPE_COLUMNS),
                      'exact_duplicate_units': 0,
                      'public_dimensions': ['province', 'district', 'affiliation', 'status', 'positivity_band']},
        }
    finally:
        wb.close()
        formula_wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--worksheet', required=True)
    parser.add_argument('--year', required=True, type=int)
    args = parser.parse_args()
    try:
        result = import_workbook(args.workbook, worksheet=args.worksheet, year=args.year)
        if args.output.exists():
            previous = json.loads(args.output.read_text(encoding='utf-8'))
            if previous.get('dataset', {}).get('version') == result['dataset']['version']:
                result['dataset']['imported_at'] = previous['dataset']['imported_at']
        args.output.parent.mkdir(parents=True, exist_ok=True)
        temporary = args.output.with_suffix('.json.tmp')
        temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        temporary.replace(args.output)
    except (ValueError, OSError, BadZipFile) as exc:
        parser.exit(1, f'นำเข้าไม่สำเร็จ: {exc}\n')
    print(f'Imported {result["audit"]["source_detail_rows"]} reporting units into {len(result["cubes"])} aggregate groups; version {result["dataset"]["version"]}.')


if __name__ == '__main__':
    main()
