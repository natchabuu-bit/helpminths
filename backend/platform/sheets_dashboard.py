"""Read-only, bounded Google Sheets statistics; source rows are never returned.

Only the supplied project sheet is public by default. Clinical denominators need
an explicit operator-approved schema; worksheet rows are not assumed to be people.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import threading
import time
import unicodedata
import zipfile
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Any

from fastapi import APIRouter, HTTPException, Query, Request

from .auth import require_user

router = APIRouter(prefix='/api/platform/operations', tags=['Google Sheets project dashboard'])
PROJECT_SHEET_ID = '1OlVnkehNCAJWVE0sR_5ypqi9XeuCI8-d'
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
THAI_DIGITS = str.maketrans('๐๑๒๓๔๕๖๗๘๙', '0123456789')
PROVINCES = {'พิษณุโลก': 'Phitsanulok', 'สุโขทัย': 'Sukhothai', 'อุตรดิตถ์': 'Uttaradit', 'เพชรบูรณ์': 'Phetchabun', 'ตาก': 'Tak'}
STATUS_LABELS = {'completed': 'ดำเนินการแล้ว', 'in_progress': 'กำลังดำเนินการ', 'pending': 'รอดำเนินการ', 'failed': 'ไม่สำเร็จ', 'cancelled': 'ยกเลิก', 'other': 'สถานะอื่น', 'unspecified': 'ไม่ระบุสถานะ'}
STATUS_ALIASES = {
    'completed': ('complete', 'completed', 'done', 'success', 'successful', 'verified', 'เสร็จสิ้น', 'เสร็จแล้ว', 'ดำเนินการแล้ว', 'สำเร็จ', 'ตรวจแล้ว', 'ยืนยันแล้ว', 'ทบทวนแล้ว', 'เรียบร้อย'),
    'in_progress': ('processing', 'inprogress', 'inreview', 'active', 'กำลังดำเนินการ', 'กำลังตรวจ', 'ระหว่างดำเนินการ'),
    'pending': ('pending', 'new', 'open', 'รอ', 'รอดำเนินการ', 'รอตรวจ', 'รอทบทวน'),
    'failed': ('failed', 'error', 'ไม่สำเร็จ', 'ล้มเหลว', 'ผิดพลาด'),
    'cancelled': ('cancelled', 'canceled', 'ยกเลิก'),
}
ALIASES = {
    'date': ('date', 'timestamp', 'created_at', 'analysis_date', 'วันที่', 'วันที่ตรวจ', 'วันที่ดำเนินการ', 'วันเดือนปี'),
    'year': ('year', 'ปี', 'ปีงบประมาณ', 'ปีพศ', 'พศ', 'ปีการศึกษา'),
    'province': ('province', 'จังหวัด', 'ชื่อจังหวัด'),
    'status': ('status', 'สถานะ', 'สถานะดำเนินการ', 'สถานะการดำเนินงาน'),
}
_CACHE_INIT_LOCK = threading.Lock()
PARASITES = {
    'Ascaris lumbricoides': ('ascarislumbricoides',),
    'Trichuris trichiura': ('trichuristrichiura',),
    'Hookworm': ('hookworm', 'hookworms', 'พยาธิปากขอ'),
    'Opisthorchis viverrini': ('opisthorchisviverrini', 'opisthorchisviverrine', 'ov'),
    'Taenia spp.': ('taeniaspp', 'taeniasp', 'taenia'),
    'Enterobius vermicularis': ('enterobiusvermicularis',),
    'Minute intestinal flukes (MIF)': ('mif', 'minuteintestinalflukes', 'พยาธิใบไม้ลำไส้ขนาดเล็ก'),
    'Hymenolepis nana': ('hymenolepisnana', 'hnana', 'พยาธิตืดแคระ', 'ตืดแคระ'),
    'Hymenolepis diminuta': ('hymenolepisdiminuta', 'hdiminuta', 'พยาธิตืดหนู', 'ตืดหนู'),
    'Echinostoma spp.': ('echinostomaspp', 'echinostoma'),
    'Fasciolopsis buski': ('fasciolopsisbuski', 'fbuski', 'พยาธิใบไม้ลำไส้ขนาดใหญ่'),
    'Strongyloides stercoralis': ('strongyloidesstercoralis',),
    # Broad source categories remain broad; a report label is not species evidence.
    'พยาธิไส้เดือน': ('พยาธิไส้เดือน', 'ไส้เดือน', 'ascaris'),
    'พยาธิแส้ม้า': ('พยาธิแส้ม้า', 'แส้ม้า', 'trichuris'),
    'พยาธิใบไม้ตับ': ('พยาธิใบไม้ตับ', 'ใบไม้ตับ'),
    'พยาธิตืด': ('พยาธิตืด', 'ตืด'),
    'พยาธิเข็มหมุด': ('พยาธิเข็มหมุด', 'เข็มหมุด', 'enterobius'),
    'พยาธิใบไม้ลำไส้ขนาดกลาง': ('พยาธิใบไม้ลำไส้ขนาดกลาง', 'ลำไส้ขนาดกลาง'),
    'สตรองจิลอย': ('สตรองจิลอย', 'สตรองจิลอยด์', 'พยาธิสตรองจิลอยด์', 'strongyloides'),
}


def _normal(value: Any) -> str:
    return re.sub(r'[\W_]+', '', str(value).translate(THAI_DIGITS).strip().lower(), flags=re.UNICODE)


def _boolean(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or not value.strip():
        return default
    if value.lower().strip() in {'1', 'true', 'yes'}:
        return True
    if value.lower().strip() in {'0', 'false', 'no'}:
        return False
    raise ValueError('Invalid boolean configuration')


@dataclass(frozen=True)
class SheetsConfig:
    enabled: bool
    spreadsheet_id: str
    gid: str
    worksheet: str
    refresh_seconds: int
    public: bool
    mapping: dict[str, Any]
    credential_json: str
    data_format: str = 'auto'

    @classmethod
    def from_environment(cls):
        try:
            sheet_id = os.getenv('HELPMINTHS_SHEETS_ID', PROJECT_SHEET_ID).strip() or PROJECT_SHEET_ID
            gid = os.getenv('HELPMINTHS_SHEETS_GID', '0').strip()
            worksheet = os.getenv('HELPMINTHS_SHEETS_WORKSHEET', '').strip()
            refresh = int(os.getenv('HELPMINTHS_SHEETS_REFRESH_SECONDS', '30'))
            mapping = json.loads(os.getenv('HELPMINTHS_SHEETS_COLUMN_MAPPING', '{}'))
            if not re.fullmatch(r'[A-Za-z0-9_-]{10,128}', sheet_id) or not re.fullmatch(r'\d{1,20}', gid):
                raise ValueError
            if len(worksheet) > 100 or any(ord(char) < 32 for char in worksheet) or not 15 <= refresh <= 300 or not isinstance(mapping, dict):
                raise ValueError
            if not isinstance(mapping.get('columns', mapping), dict) or type(mapping.get('header_row', 1)) is not int or not 1 <= mapping.get('header_row', 1) <= 100:
                raise ValueError
            credential = os.getenv('HELPMINTHS_GOOGLE_SERVICE_ACCOUNT_JSON', '').strip()
            data_format = os.getenv('HELPMINTHS_SHEETS_FORMAT', 'auto').strip().lower()
            if len(credential) > 65536 or data_format not in {'auto', 'csv', 'xlsx'}:
                raise ValueError
            return cls(_boolean('HELPMINTHS_SHEETS_ENABLED', True), sheet_id, gid, worksheet, refresh,
                       _boolean('HELPMINTHS_SHEETS_PUBLIC_AGGREGATES', sheet_id == PROJECT_SHEET_ID), mapping, credential, data_format)
        except (ValueError, TypeError):
            raise HTTPException(503, 'ตั้งค่าตัวเชื่อม Google Sheets ไม่ถูกต้อง กรุณาตรวจค่าบนเซิร์ฟเวอร์') from None

    def source(self):
        return {'provider': 'google_sheets', 'spreadsheet_id': self.spreadsheet_id,
                'url': f'https://docs.google.com/spreadsheets/d/{self.spreadsheet_id}/edit',
                'gid': self.gid, 'worksheet': self.worksheet or None, 'format': self.data_format}

    def metadata(self):
        return {'configured': self.enabled, 'source': self.source(), 'refresh_seconds': self.refresh_seconds,
                'access': 'public_aggregates' if self.public else 'authenticated',
                'mapping_configured': bool(self.mapping), 'capabilities': {'read_only': True, 'raw_rows_exposed': False,
                    'filters': ['year', 'province', 'from_date', 'to_date', 'status'],
                    'person_counts_require_explicit_schema': True, 'formats': ['csv', 'xlsx'],
                    'authentication': 'service_account' if self.credential_json else 'public_readonly'}}


class SourceError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


class SafeRedirects(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, message, headers, newurl):
        target = urllib.parse.urlsplit(newurl)
        allowed = target.hostname in {'docs.google.com', 'sheets.googleapis.com', 'drive.usercontent.google.com', 'drive.google.com'} or (target.hostname or '').endswith('.googleusercontent.com')
        if request.has_header('Authorization') or target.scheme != 'https' or not allowed:
            raise SourceError('access_denied', 'Google Sheet ยังอ่านไม่ได้ กรุณาตรวจสิทธิ์การแชร์หรือบัญชีบริการ')
        return super().redirect_request(request, fp, code, message, headers, newurl)


def fetch_bytes(url: str, token: str | None = None) -> tuple[bytes, str]:
    headers = {'User-Agent': 'HelpMinths-ReadOnly-Dashboard/1.0',
               'Accept': 'text/csv,application/json,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream'}
    if token:
        headers['Authorization'] = f'Bearer {token}'
    try:
        with urllib.request.build_opener(SafeRedirects()).open(urllib.request.Request(url, headers=headers), timeout=12) as response:
            body = response.read(MAX_BYTES + 1)
            if len(body) > MAX_BYTES:
                raise SourceError('source_too_large', 'ชีตมีข้อมูลเกินขีดจำกัด 5 MB กรุณาใช้แท็บสรุปขนาดเล็ก')
            content_type = response.headers.get('Content-Type', '').lower()
            if 'html' in content_type or body.lstrip().lower().startswith((b'<!doctype html', b'<html')):
                raise SourceError('access_denied', 'Google ส่งหน้าลงชื่อเข้าใช้แทนข้อมูล กรุณาตรวจสิทธิ์ชีต')
            return body, content_type
    except urllib.error.HTTPError as error:
        if error.code in {401, 403}:
            raise SourceError('access_denied', 'อ่านชีตไม่ได้ กรุณาตรวจสิทธิ์การแชร์หรือบัญชีบริการ') from None
        if error.code == 404:
            raise SourceError('source_not_found', 'ไม่พบ Google Sheet หรือแท็บที่กำหนด กรุณาตรวจลิงก์และสิทธิ์') from None
        if error.code == 400:
            raise SourceError('invalid_sheet_data', 'Google อ่านตารางนี้ไม่ได้ กรุณาตรวจชื่อแท็บและชนิดไฟล์ Google Sheets') from None
        if error.code == 429:
            raise SourceError('quota_exceeded', 'Google จำกัดคำขอชั่วคราว ระบบจะลองใหม่ตามรอบอัปเดต') from None
        raise SourceError('source_unavailable', 'Google Sheets ไม่พร้อมตอบข้อมูลชั่วคราว') from None
    except (urllib.error.URLError, TimeoutError, OSError):
        raise SourceError('connection_failed', 'เชื่อมต่อ Google Sheets ไม่สำเร็จ อาจเกิดจากเครือข่ายหรือสิทธิ์ของแหล่งข้อมูล') from None


def read_excel(body: bytes, config: SheetsConfig) -> list[list[Any]]:
    """Read cached cell values only; never evaluate formulas or run Office macros."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SourceError('dependency_not_configured', 'ตัวอ่าน Excel ต้องติดตั้ง backend/requirements.lock ก่อน') from None
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(item.file_size for item in items) > 50 * 1024 * 1024 or any(item.flag_bits & 1 for item in items):
                raise SourceError('source_too_large', 'ไฟล์ Excel ขยายใหญ่เกินกำหนดหรือมีการเข้ารหัส กรุณาใช้แท็บสรุปขนาดเล็ก')
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_links=False)
        try:
            if config.worksheet:
                if config.worksheet not in workbook.sheetnames:
                    raise SourceError('source_not_found', 'ไม่พบชื่อแท็บที่กำหนดในไฟล์ Excel')
                worksheet = workbook[config.worksheet]
            else:
                worksheet = workbook.worksheets[0]
            if (worksheet.max_row or 0) > MAX_ROWS + 104 or (worksheet.max_column or 0) > 200:
                raise SourceError('source_too_large', 'ไฟล์ Excel เกินขีดจำกัดแถวหรือคอลัมน์ กรุณาใช้แท็บสรุปขนาดเล็ก')
            table = []
            for row in worksheet.iter_rows(values_only=True):
                if len(table) >= MAX_ROWS + 104 or len(row) > 200:
                    raise SourceError('source_too_large', 'ไฟล์ Excel เกินขีดจำกัดแถวหรือคอลัมน์ กรุณาใช้แท็บสรุปขนาดเล็ก')
                table.append([value if value is not None else '' for value in row])
            return table
        finally:
            workbook.close()
    except SourceError:
        raise
    except Exception:
        raise SourceError('invalid_sheet_data', 'อ่านไฟล์ Excel ไม่สำเร็จ กรุณาใช้ .xlsx ที่สมบูรณ์และไม่เข้ารหัส') from None


def _public_excel(config: SheetsConfig) -> list[list[Any]]:
    parameters = urllib.parse.urlencode({'id': config.spreadsheet_id, 'export': 'download'})
    body, _ = fetch_bytes(f'https://drive.usercontent.google.com/download?{parameters}')
    return read_excel(body, config)


def _fetch_table(config: SheetsConfig) -> list[list[Any]]:
    if config.credential_json:
        try:
            from google.auth.transport.requests import Request as GoogleRequest
            from google.oauth2 import service_account
            info = json.loads(config.credential_json)
            if info.get('token_uri') not in {'https://oauth2.googleapis.com/token', 'https://www.googleapis.com/oauth2/v4/token'}:
                raise ValueError
            scopes = ['https://www.googleapis.com/auth/drive.readonly'] if config.data_format == 'xlsx' else ['https://www.googleapis.com/auth/spreadsheets.readonly']
            credentials = service_account.Credentials.from_service_account_info(info, scopes=scopes)
            transport = GoogleRequest()
            deadline = time.monotonic() + 12
            def bounded_request(*args, **kwargs):
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise SourceError('connection_failed', 'การยืนยันบัญชีบริการใช้เวลานานเกินไป ระบบจะลองใหม่ตามรอบ')
                kwargs['timeout'] = min(12, remaining)
                return transport(*args, **kwargs)
            credentials.refresh(bounded_request)
            if config.data_format == 'xlsx':
                body, _ = fetch_bytes(f'https://www.googleapis.com/drive/v3/files/{config.spreadsheet_id}?alt=media', credentials.token)
                return read_excel(body, config)
            worksheet = config.worksheet
            if not worksheet:
                # Read the actual first tab name instead of assuming "Sheet1"
                # or requesting a rectangular range beyond its grid limits.
                metadata, _ = fetch_bytes(f'https://sheets.googleapis.com/v4/spreadsheets/{config.spreadsheet_id}?fields=sheets.properties.title', credentials.token)
                tabs = json.loads(metadata).get('sheets', [])
                worksheet = tabs[0]['properties']['title'] if tabs else None
                if not isinstance(worksheet, str) or not worksheet:
                    raise SourceError('source_not_found', 'ไม่พบแท็บที่อ่านได้ใน Google Sheet')
            cell_range = "'" + worksheet.replace("'", "''") + "'"
            url = f'https://sheets.googleapis.com/v4/spreadsheets/{config.spreadsheet_id}/values/{urllib.parse.quote(cell_range, safe="")}'
            body, _ = fetch_bytes(url, credentials.token)
            payload = json.loads(body)
            table = payload.get('values')
            if not isinstance(table, list):
                raise ValueError
            return table
        except ImportError:
            raise SourceError('dependency_not_configured', 'ตัวอ่านชีตส่วนตัวต้องติดตั้ง backend/requirements-sheets.txt ก่อน') from None
        except SourceError:
            raise
        except Exception:
            raise SourceError('credentials_unavailable', 'บัญชีบริการอ่านชีตไม่ได้ กรุณาตรวจ credentials และสิทธิ์บนเซิร์ฟเวอร์') from None
    if config.data_format == 'xlsx':
        return _public_excel(config)
    parameters = {'format': 'csv', 'gid': config.gid}
    if config.worksheet:
        # gviz supports a worksheet name without publishing a new URL or changing sharing.
        parameters = {'tqx': 'out:csv', 'sheet': config.worksheet, 'headers': '0'}
        url = f'https://docs.google.com/spreadsheets/d/{config.spreadsheet_id}/gviz/tq?{urllib.parse.urlencode(parameters)}'
    else:
        url = f'https://docs.google.com/spreadsheets/d/{config.spreadsheet_id}/export?{urllib.parse.urlencode(parameters)}'
    body, content_type = fetch_bytes(url)
    if 'json' in content_type or body.lstrip().startswith((b'{', b'[')):
        raise SourceError('invalid_sheet_data', 'Google ไม่ได้ส่งตาราง CSV กรุณาตรวจชนิดไฟล์ แท็บ และสิทธิ์การอ่าน')
    try:
        text = body.decode('utf-8-sig')
        return list(csv.reader(io.StringIO(text), strict=True))
    except (UnicodeError, csv.Error):
        raise SourceError('invalid_sheet_data', 'ข้อมูลชีตไม่ใช่ CSV ที่อ่านได้ กรุณาใช้ Google Sheet หรือแท็บสรุปที่รองรับ') from None


def fetch_table(config: SheetsConfig) -> list[list[Any]]:
    try:
        return _fetch_table(config)
    except SourceError as error:
        # Office-editing Google URLs can point to an XLSX, not a native Sheet.
        # A failed format export may be retried as a read of the same file ID.
        if config.data_format == 'auto' and error.code in {'invalid_sheet_data', 'access_denied', 'source_not_found'}:
            return _fetch_table(SheetsConfig(config.enabled, config.spreadsheet_id, config.gid, config.worksheet,
                                           config.refresh_seconds, config.public, config.mapping, config.credential_json, 'xlsx'))
        raise


def _date(value: Any) -> date | None:
    text = str(value).translate(THAI_DIGITS).strip()
    if not text:
        return None
    text = text.split('T')[0].split(' ')[0]
    try:
        parts = re.split(r'[-/]', text)
        if len(parts) != 3:
            return None
        first, second, third = map(int, parts)
        year, month, day = (first, second, third) if first > 31 else (third, second, first)
        if year >= 2400:
            year -= 543
        return date(year, month, day)
    except (ValueError, OverflowError):
        return None


def _year(value: Any) -> int | None:
    text = str(value).translate(THAI_DIGITS).strip().replace('พ.ศ.', '').replace('พศ', '').strip()
    try:
        number = int(text)
        if number >= 2400:
            number -= 543
        return number + 543 if 1900 <= number <= 2200 else None
    except (ValueError, TypeError):
        return None


def _count(value: Any, person: bool = False, infected: bool = False) -> int | None:
    text = str(value).translate(THAI_DIGITS).strip()
    if person:
        normalized = _normal(text)
        yes = {'1', 'true', 'yes', 'positive', 'พบ', 'พบเชื้อ', 'ตรวจพบ', 'ติดเชื้อ'} if infected else {'1', 'true', 'yes', 'tested', 'ตรวจแล้ว', 'ได้รับการตรวจ'}
        no = {'0', 'false', 'no', 'negative', 'ไม่พบ', 'ไม่พบเชื้อ', 'ไม่ติดเชื้อ'} if infected else {'0', 'false', 'no', 'untested', 'ยังไม่ตรวจ', 'ไม่ได้ตรวจ'}
        return 1 if normalized in {_normal(value) for value in yes} else 0 if normalized in {_normal(value) for value in no} else None
    if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.0+)?', text):
        return None
    value = float(text.replace(',', ''))
    return int(value) if math.isfinite(value) and 0 <= value <= 10**12 else None


SCHOOL_PARASITES = [(9, 'พยาธิปากขอ'), (11, 'พยาธิไส้เดือน'), (13, 'พยาธิแส้ม้า'), (15, 'พยาธิใบไม้ตับ'),
                    (17, 'พยาธิเข็มหมุด'), (19, 'พยาธิตืด'), (21, 'พยาธิใบไม้ลำไส้ขนาดกลาง'), (23, 'สตรองจิลอย')]


def _school_total(value: Any) -> bool:
    normalized = _normal(value)
    if normalized in {_normal(label) for label in ('รวม', 'รวมทั้งสิ้น', 'ยอดรวม', 'รวมทั้งหมด', 'total', 'subtotal', 'grand total')}:
        return True
    if normalized.startswith(tuple(_normal(label) for label in ('รวมจังหวัด', 'รวมอำเภอ', 'รวมตำบล', 'รวมทั้งจังหวัด', 'รวมทั้งอำเภอ', 'รวมทั้งตำบล', 'รวมทั้งโครงการ'))):
        return True
    prefixes = '|'.join(_normal(label) for label in ('รวม', 'ยอดรวม', 'รวมจำนวน', 'total', 'subtotal', 'grand total'))
    units = '|'.join(_normal(label) for label in ('โรงเรียน', 'แห่ง', 'school', 'schools'))
    # Match a subtotal's count, while preserving real names such as รวมไทยพัฒนา1.
    return bool(re.fullmatch(f'(?:{prefixes})[0-9]+(?:{units})', normalized))


def _school_identity(value: Any) -> str:
    # Thai vowels/tone marks distinguish institutions; unlike header matching,
    # identity must not remove Unicode combining marks or other name characters.
    return ' '.join(unicodedata.normalize('NFC', str(value)).translate(THAI_DIGITS).casefold().split())


def _province(value: Any) -> str | None:
    normalized = _normal(value).removeprefix('จังหวัด').removeprefix('จ')
    return next((name for name, english in PROVINCES.items() if normalized in {_normal(name), _normal(english)}), None)


def _school_header(table: list[list[Any]]) -> int | None:
    """Match the full user-provided three-row header, never data by position alone."""
    for start in range(min(100, max(0, len(table) - 2))):
        header = table[start:start + 3]
        if any(not isinstance(row, list) for row in header):
            continue
        combined = [''.join(_normal(row[index]) for row in header if index < len(row)) for index in range(28)]
        expected = {0: 'โรงเรียน', 1: 'ตำบล', 2: 'อำเภอ', 3: 'จังหวัด', 4: 'นักเรียนทั้งหมด',
                    5: 'นักเรียนส่งตรวจ', 6: 'ครอบคลุม', 7: 'ตรวจพบพยาธิ', 8: 'ร้อยละ', 9: 'ปากขอ',
                    11: 'ไส้เดือน', 13: 'แส้ม้า', 15: 'ใบไม้ตับ', 17: 'เข็มหมุด', 19: 'ตืด',
                    21: 'ลำไส้ขนาดกลาง', 23: 'สตรองจิลอย', 25: 'มากกว่า1ชนิด', 27: 'สังกัด'}
        if not all(_normal(label) in combined[index] for index, label in expected.items()):
            continue
        if not all(_normal('จำนวน') in combined[index] and _normal('ร้อยละ') in combined[index + 1] for index in range(9, 26, 2)):
            continue
        return start + 3
    return None


def _school_rows(table: list[list[Any]], start: int) -> tuple[list[dict], dict, list[str]]:
    if len(table) - start > MAX_ROWS or any(len(row) > 200 for row in table):
        raise SourceError('source_too_large', 'แท็บสรุปโรงเรียนเกินขีดจำกัดที่ระบบรองรับ')
    rows, identities, warnings = [], {}, []
    skipped = duplicates = conflicts = invalid_counts = missing_province = 0
    for source in table[start:]:
        if not any(str(cell).strip() for cell in source):
            continue
        cell = lambda index: source[index] if index < len(source) else ''
        school = _normal(cell(0))
        if not school or _school_total(cell(0)):
            skipped += 1
            continue
        province = _province(cell(3))
        missing_province += int(province is None)
        identity = hashlib.sha256(json.dumps([_school_identity(cell(index)) for index in (0, 1, 2, 3)], ensure_ascii=False).encode()).hexdigest()
        total, tested, infected = _count(cell(4)), _count(cell(5)), _count(cell(7))
        inconsistent = (total is not None and tested is not None and tested > total) or (tested is not None and infected is not None and infected > tested)
        if inconsistent:
            total = tested = infected = None
            invalid_counts += 1
        parasites = {label: _count(cell(index)) for index, label in SCHOOL_PARASITES}
        for label, count in parasites.items():
            if count is not None and infected is not None and count > infected:
                parasites[label] = None
                invalid_counts += 1
        multiple = _count(cell(25))
        if multiple is not None and infected is not None and multiple > infected:
            multiple = None
            invalid_counts += 1
        row = {'date': None, 'year': None, 'province': province, 'status': 'unspecified', 'tested': tested,
               'infected': infected, 'total_students': total, 'multiple_infections': multiple,
               'school_id': identity, 'school_parasites': parasites, 'parasite': None, 'count': None}
        if identity in identities:
            previous = identities[identity]
            if previous == row:
                duplicates += 1
                continue
            conflicts += 1
            previous.update(tested=None, infected=None, total_students=None, multiple_infections=None,
                            school_parasites={label: None for _, label in SCHOOL_PARASITES})
            continue
        identities[identity] = row
        rows.append(row)
    if skipped:
        warnings.append(f'ละเว้นแถวสรุปรวมหรือแถวที่ไม่มีโรงเรียน {skipped} แถว เพื่อไม่บวกยอดซ้ำ')
    if duplicates:
        warnings.append(f'พบแถวโรงเรียนซ้ำที่มีตัวเลขตรงกัน {duplicates} แถว จึงนับเพียงครั้งเดียว')
    if conflicts:
        warnings.append(f'พบข้อมูลโรงเรียนซ้ำแต่ตัวเลขไม่ตรงกัน {conflicts} แถว ต้องตรวจทานก่อนใช้ยอดรวม')
    if invalid_counts:
        warnings.append(f'พบตัวเลขขัดกัน {invalid_counts} จุด เช่น ส่งตรวจเกินนักเรียน หรือผลบวกเกินส่งตรวจ จึงไม่สรุปช่องนั้น')
    if missing_province:
        warnings.append(f'ข้อมูลต้นทางมี {missing_province} โรงเรียนที่ไม่ระบุจังหวัดในเขตสุขภาพที่ 2 ยอดรวมทั้งหมดจึงยังไม่พร้อม; ตัวกรองจังหวัดใช้เฉพาะแถวที่ระบุจังหวัดนั้นได้')
    if any(row['total_students'] is None or row['tested'] is None or row['infected'] is None or row['multiple_infections'] is None or any(value is None for value in row['school_parasites'].values()) for row in rows):
        warnings.append('บางช่องจำนวนไม่มีข้อมูลที่ใช้คำนวณได้ จึงแสดงช่องนั้นเป็นไม่พร้อมแทนการตีความช่องว่างหรือขีดเป็นศูนย์')
    return rows, {'matched_roles': ['province', 'tested', 'infected', 'total_students', 'school', 'school_parasites', 'multiple_infections'],
                  'schema_valid': True, 'schema_id': 'health2-kpd-school-v1', 'province_mapped': True,
                  'status_mapped': False, 'has_unknown_area': bool(missing_province)}, warnings


def parse_table(table: list[list[Any]], config: SheetsConfig) -> tuple[list[dict], dict, list[str]]:
    if not table:
        return [], {'matched_roles': [], 'schema_valid': False, 'status_mapped': False}, []
    if not isinstance(table, list) or any(not isinstance(row, list) or len(row) > 200 for row in table):
        raise SourceError('invalid_sheet_data', 'ตารางต้นทางมีโครงสร้างไม่ถูกต้องหรือเกิน 200 คอลัมน์')
    if (config.spreadsheet_id == PROJECT_SHEET_ID and not config.mapping) or config.mapping.get('preset') == 'health2-kpd-school-v1':
        school_start = _school_header(table)
        if school_start is not None:
            return _school_rows(table, school_start)
    header_index = config.mapping.get('header_row', 1) - 1
    if len(table) > MAX_ROWS + header_index + 1 or any(not isinstance(row, list) or len(row) > 200 for row in table):
        raise SourceError('source_too_large', 'ชีตเกินขีดจำกัด 10,000 แถวหรือ 200 คอลัมน์ กรุณาใช้แท็บสรุป')
    if header_index >= len(table):
        raise SourceError('mapping_required', 'แถวหัวตารางที่กำหนดไม่อยู่ในข้อมูลที่อ่านได้')
    headers = [_normal(cell) for cell in table[header_index]]
    mapping = config.mapping.get('columns', config.mapping)
    roles, warnings = {}, []
    for role in ('date', 'year', 'province', 'status', 'tested', 'infected', 'parasite', 'count'):
        requested = mapping.get(role)
        if requested is not None and not isinstance(requested, str):
            raise SourceError('mapping_required', 'ชื่อคอลัมน์ที่กำหนดต้องเป็นข้อความ')
        names = [_normal(requested)] if requested else [_normal(value) for value in ALIASES.get(role, ())]
        found = [index for index, header in enumerate(headers) if header and header in names]
        if len(found) == 1:
            roles[role] = found[0]
        elif requested or len(found) > 1:
            warnings.append(f'ต้องตรวจการจับคู่คอลัมน์ {role} บนเซิร์ฟเวอร์')
    schema_valid = config.mapping.get('row_granularity') in {'person', 'aggregate'} and config.mapping.get('rows_disjoint') is True and config.mapping.get('population_unit') == 'people'
    person = config.mapping.get('row_granularity') == 'person'
    rows = []
    for source in table[header_index + 1:]:
        if not any(str(cell).strip() for cell in source):
            continue
        def value(role):
            index = roles.get(role)
            return source[index] if index is not None and index < len(source) else ''
        canonical_province = _province(value('province'))
        parsed_date = _date(value('date'))
        parasite = next((name for name, aliases in PARASITES.items() if _normal(value('parasite')) in {_normal(name), *(_normal(alias) for alias in aliases)}), None)
        raw_status = _normal(value('status'))
        status = next((key for key, aliases in STATUS_ALIASES.items() if raw_status in {_normal(alias) for alias in aliases}), 'other' if raw_status else 'unspecified')
        rows.append({'date': parsed_date.isoformat() if parsed_date else None,
                     'year': _year(value('year')) or (parsed_date.year + 543 if parsed_date else None),
                     'province': canonical_province, 'status': status,
                     'tested': _count(value('tested'), person=person) if 'tested' in roles else None,
                     'infected': _count(value('infected'), person=person, infected=True) if 'infected' in roles else None,
                     'parasite': parasite, 'count': _count(value('count')) if 'count' in roles else None})
    if not schema_valid:
        warnings.append('ยังไม่ยืนยันโครงสร้างข้อมูลและการไม่ซ้ำของบุคคล จึงไม่สรุปจำนวนผู้ตรวจ ผู้ติดเชื้อ หรือสัดส่วนผลบวกจากจำนวนแถว')
    if 'province' in roles and any(row['province'] is None for row in rows):
        warnings.append('บางแถวไม่ระบุจังหวัดในเขตสุขภาพที่ 2 หรือเป็นแถวสรุปรวม จึงไม่ใช้แถวนั้นคำนวณยอดรายบุคคล')
    if 'date' in roles and any(row['date'] is None for row in rows):
        warnings.append('บางแถวไม่มีวันที่ที่อ่านได้ จึงไม่รวมในแนวโน้มรายวันหรือช่วงวันที่')
    return rows, {'matched_roles': sorted(roles), 'schema_valid': schema_valid, 'status_mapped': 'status' in roles, 'province_mapped': 'province' in roles}, warnings


def aggregate(rows: list[dict], info: dict, config: SheetsConfig, filters: dict) -> tuple[dict, list[str]]:
    warnings = []
    all_rows = rows
    for key in ('year', 'province', 'status'):
        if filters.get(key) is not None:
            rows = [row for row in rows if row[key] == filters[key]]
    if filters.get('from_date'):
        rows = [row for row in rows if row['date'] is not None and row['date'] >= filters['from_date']]
    if filters.get('to_date'):
        rows = [row for row in rows if row['date'] is not None and row['date'] <= filters['to_date']]
    # Exclude non-area and total rows from person counts whenever a province field exists.
    count_rows = [row for row in rows if row['province'] is not None] if info.get('province_mapped') else rows
    def people_counts(group):
        valid = info['schema_valid'] and bool(group)
        tested = sum(row['tested'] for row in group) if valid and all(row['tested'] is not None for row in group) else None
        infected = sum(row['infected'] for row in group) if valid and all(row['infected'] is not None for row in group) else None
        if tested is not None and infected is not None and (infected > tested or any(row['infected'] > row['tested'] for row in group)):
            return None, None
        return tested, infected
    tested, infected = people_counts(count_rows)
    unknown_area = info.get('schema_id') == 'health2-kpd-school-v1' and any(row['province'] is None for row in rows)
    if unknown_area:
        tested = infected = None
    rate = round(infected / tested * 100, 2) if tested and infected is not None else None
    if info['schema_valid'] and count_rows and (tested is None or infected is None):
        warnings.append('ข้อมูลจำนวนคนบางแถวขาดหาย ไม่ถูกต้อง หรือจำนวนผลบวกเกินจำนวนผู้ตรวจ จึงไม่แสดงยอดรวมที่อาจคลาดเคลื่อน')
    province_groups = defaultdict(list)
    for row in rows:
        if row['province']:
            province_groups[row['province']].append(row)
    provinces = []
    for province, group in sorted(province_groups.items()):
        group_tested, group_infected = people_counts(group)
        entry = {'province': province, 'records': len(group), 'tested_count': group_tested, 'infected_count': group_infected}
        if info.get('schema_id') == 'health2-kpd-school-v1':
            entry['school_count'] = len(group)
            entry['total_students'] = sum(row['total_students'] for row in group) if all(row['total_students'] is not None for row in group) else None
        provinces.append(entry)
    by_parasite = []
    total_students = coverage = multiple = school_count = None
    if info.get('schema_id') == 'health2-kpd-school-v1':
        school_count = len(rows)
        eligible = count_rows and not unknown_area
        total_students = sum(row['total_students'] for row in count_rows) if eligible and all(row['total_students'] is not None for row in count_rows) else None
        multiple = sum(row['multiple_infections'] for row in count_rows) if eligible and all(row['multiple_infections'] is not None for row in count_rows) else None
        coverage = round(tested / total_students * 100, 2) if total_students and tested is not None else None
        for _, label in SCHOOL_PARASITES:
            if eligible and all(row['school_parasites'][label] is not None for row in count_rows):
                by_parasite.append({'parasite': label, 'count': sum(row['school_parasites'][label] for row in count_rows), 'unit': 'people'})
        if by_parasite:
            warnings.append('จำนวนแยกชนิดอาจมีการติดเชื้อร่วม ไม่นำมาบวกแทนจำนวนผู้ตรวจพบพยาธิทั้งหมด; มากกว่า 1 ชนิดแสดงแยกต่างหาก')
    unit = config.mapping.get('parasite_counts_unit')
    if config.mapping.get('rows_disjoint') is True and unit in {'people', 'objects', 'positive_tests'} and 'parasite' in info['matched_roles'] and 'count' in info['matched_roles']:
        parasite_groups = defaultdict(list)
        for row in count_rows:
            if row['parasite']:
                parasite_groups[row['parasite']].append(row)
        for name, group in sorted(parasite_groups.items()):
            if all(row['count'] is not None for row in group):
                by_parasite.append({'parasite': name, 'count': sum(row['count'] for row in group), 'unit': unit})
        if by_parasite:
            warnings.append('จำนวนแยกชนิดพยาธิอาจมีการติดเชื้อร่วม ห้ามบวกเป็นจำนวนผู้ติดเชื้อไม่ซ้ำทั้งหมด')
    trend = Counter(row['date'] for row in rows if row['date'])
    statuses = Counter(row['status'] for row in rows) if info['status_mapped'] else {}
    metrics = [{'key': 'total_students', 'label': 'นักเรียนทั้งหมด', 'value': total_students, 'unit': 'คน'},
               {'key': 'coverage_percent', 'label': 'ความครอบคลุมการส่งตรวจ', 'value': coverage, 'unit': '%'},
               {'key': 'multiple_infections_count', 'label': 'ตรวจพบมากกว่า 1 ชนิด', 'value': multiple, 'unit': 'คน'},
               {'key': 'school_count', 'label': 'โรงเรียนในข้อมูลที่เลือก', 'value': school_count, 'unit': 'โรงเรียน'},
               {'key': 'row_count', 'label': 'แถวข้อมูลที่เลือก', 'value': len(rows), 'unit': 'แถว'},
               {'key': 'tested_count', 'label': 'นักเรียนส่งตรวจ', 'value': tested, 'unit': 'คน'},
               {'key': 'infected_count', 'label': 'ผู้ตรวจพบเชื้อ', 'value': infected, 'unit': 'คน'},
               {'key': 'infection_rate_percent', 'label': 'ร้อยละตรวจพบในกลุ่มส่งตรวจ', 'value': rate, 'unit': '%'},
               {'key': 'province_count', 'label': 'จังหวัดที่มีในข้อมูลที่เลือก', 'value': len(provinces) if info.get('province_mapped') else None, 'unit': 'จังหวัด'}]
    return {'row_count': len(rows), 'schema_id': info.get('schema_id'), 'total_students': total_students, 'coverage_percent': coverage,
            'multiple_infections_count': multiple, 'school_count': school_count, 'tested_count': tested, 'infected_count': infected, 'infection_rate_percent': rate,
            'metrics': metrics, 'trend': [{'date': day, 'count': count} for day, count in sorted(trend.items())],
            'statuses': [{'key': key, 'label': STATUS_LABELS[key], 'count': count} for key, count in sorted(statuses.items())],
            'by_province': provinces, 'by_parasite': by_parasite,
            'available_years': sorted({row['year'] for row in all_rows if row['year']}),
            'available_provinces': sorted({row['province'] for row in all_rows if row['province']}),
            'filters_applied': filters, 'coverage': {'dated_rows': sum(row['date'] is not None for row in rows),
                'status_rows': sum(row['status'] != 'unspecified' for row in rows) if info['status_mapped'] else 0}}, warnings


class SheetsCache:
    def __init__(self):
        self.lock = threading.Lock()
        self.rows = None
        self.info = None
        self.warnings = []
        self.fetched_at = None
        self.last_attempt_at = None
        self.next_attempt = 0.0
        self.error = None

    def snapshot(self, config: SheetsConfig, filters: dict):
        with self.lock:
            if time.monotonic() >= self.next_attempt:
                self.last_attempt_at = datetime.now(timezone.utc).isoformat()
                try:
                    self.rows, self.info, self.warnings = parse_table(fetch_table(config), config)
                    self.fetched_at = datetime.now(timezone.utc).isoformat()
                    self.error = None
                except SourceError as error:
                    self.error = {'code': error.code, 'message': error.message}
                except Exception:
                    self.error = {'code': 'source_unavailable', 'message': 'อ่านข้อมูล Google Sheets ไม่สำเร็จ กรุณาตรวจการตั้งค่าและลองใหม่ตามรอบ'}
                self.next_attempt = time.monotonic() + config.refresh_seconds
            base = {'source': config.source(), 'refresh_seconds': config.refresh_seconds, 'fetched_at': self.fetched_at,
                    'last_attempt_at': self.last_attempt_at, 'error': self.error, 'warnings': list(self.warnings), 'data': None}
            if self.rows is None:
                return {**base, 'status': 'unavailable'}
            data, warnings = aggregate(self.rows, self.info, config, filters)
            if self.error:
                status = 'stale'
            elif not self.rows:
                status = 'empty'
            elif not self.info['schema_valid']:
                status = 'mapping_required'
            else:
                status = 'ready'
            return {**base, 'status': status, 'warnings': [*self.warnings, *warnings], 'data': data}


@router.get('/status')
def operations_status():
    return SheetsConfig.from_environment().metadata()


@router.get('/dashboard')
def operations_dashboard(request: Request, year: int | None = Query(None, ge=1900, le=2743),
                         province: str | None = Query(None, max_length=60), from_date: str | None = Query(None, max_length=10),
                         to_date: str | None = Query(None, max_length=10), status: str | None = Query(None, max_length=20)):
    config = SheetsConfig.from_environment()
    if not config.public:
        require_user(request)
    if not config.enabled:
        return {'status': 'not_configured', 'source': config.source(), 'refresh_seconds': config.refresh_seconds,
                'fetched_at': None, 'last_attempt_at': None, 'error': None, 'warnings': [], 'data': None}
    filters = {}
    if year is not None:
        selected_year = _year(year)
        if selected_year is None:
            raise HTTPException(422, 'ปีต้องเป็น ค.ศ. 1900–2200 หรือ พ.ศ. ที่สอดคล้องกัน')
        filters['year'] = selected_year
    if province:
        if province not in PROVINCES:
            raise HTTPException(422, 'เลือกจังหวัดในเขตสุขภาพที่ 2 จากรายการที่ระบบรองรับ')
        filters['province'] = province
    for key, value in (('from_date', from_date), ('to_date', to_date)):
        if value:
            parsed = _date(value)
            if parsed is None:
                raise HTTPException(422, 'วันที่ต้องเป็น YYYY-MM-DD')
            filters[key] = parsed.isoformat()
    if filters.get('from_date', '') > filters.get('to_date', '9999-12-31'):
        raise HTTPException(422, 'วันเริ่มต้นต้องไม่อยู่หลังวันสิ้นสุด')
    if status:
        if status not in STATUS_LABELS:
            raise HTTPException(422, 'สถานะไม่ถูกต้อง')
        filters['status'] = status
    # Configuration changes never reuse another source/permission/schema cache.
    signature = (config.spreadsheet_id, config.gid, config.worksheet, config.refresh_seconds, config.public,
                 json.dumps(config.mapping, sort_keys=True), config.credential_json, config.data_format)
    with _CACHE_INIT_LOCK:
        if not hasattr(request.app.state, 'sheets_caches'):
            request.app.state.sheets_caches = {}
            request.app.state.sheets_caches_lock = threading.Lock()
    with request.app.state.sheets_caches_lock:
        cache = request.app.state.sheets_caches.get(signature)
        if cache is None:
            # Keep only the current configuration: source data are not an archive.
            request.app.state.sheets_caches = {signature: (cache := SheetsCache())}
    return cache.snapshot(config, filters)
