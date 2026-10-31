"""Synthetic workbook fixtures verify the importer; these are not health evidence."""
from __future__ import annotations

import copy
import json
import sys
from zipfile import BadZipFile

import pytest
from openpyxl import Workbook, load_workbook

from backend.platform import import_survey as importer

STAMP = '2026-10-30T12:00:00Z'


def record(name='PRIVATE_SCHOOL_ALPHA', *, affiliation='สังกัด A', enrolled=10, tested=8, infected=2, coinfection=0):
    types = {key: 0 for key, _, _ in importer.TYPE_COLUMNS}
    if infected:
        types['ascaris'] = infected
    return {'name': name, 'tambon': 'ตำบล A', 'district': 'อำเภอ A', 'province': 'จังหวัด A',
            'affiliation': affiliation, 'enrolled': enrolled, 'tested': tested, 'infected': infected,
            'coinfection': coinfection, 'types': types}


def example_records():
    first = record()
    zero = record('PRIVATE_SCHOOL_ZERO', enrolled=0, tested=0, infected=0)
    zero['types'].update(ascaris=None, pinworm=1)
    pending = record('PRIVATE_CHILD_CENTER_PENDING', enrolled=None, tested=None, infected=None, coinfection=None)
    pending['types'] = {key: None for key in pending['types']}
    negative = record('PRIVATE_SCHOOL_NEGATIVE', enrolled=5, tested=5, infected=0)
    negative['types']['tapeworm'] = None
    return [first, zero, pending, negative]


def build_workbook(path, records=None, *, leading=True, sheet_name='เขต2', extra_summaries=True):
    records = copy.deepcopy(records if records is not None else example_records())
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    start = 2 if leading else 1  # openpyxl uses one-based columns
    if leading:
        sheet.cell(2, start - 1, 'ที่')
    for offset, label in {0: 'โรงเรียน', 1: 'ตำบล', 2: 'อำเภอ', 3: 'จังหวัด',
                          4: 'จำนวนนักเรียนทั้งหมด', 5: 'จำนวนนักเรียนส่งตรวจ',
                          6: 'ความครอบคลุม', 7: 'ตรวจพบพยาธิ', 27: 'สังกัด'}.items():
        sheet.cell(2, start + offset, label)
    for _, label, offset in importer.TYPE_COLUMNS:
        sheet.cell(3, start + offset, label)
        sheet.cell(4, start + offset, 'จำนวน (คน)')
    sheet.cell(3, start + 25, 'มากกว่า 1 ชนิด')
    for index, item in enumerate(records, start=5):
        if leading:
            sheet.cell(index, start - 1, index - 4)
        for offset, field in {0: 'name', 1: 'tambon', 2: 'district', 3: 'province', 27: 'affiliation'}.items():
            sheet.cell(index, start + offset, item[field])
        for field, offset in importer.CORE_COLUMNS.items():
            sheet.cell(index, start + offset, item[field])
        for field, _, offset in importer.TYPE_COLUMNS:
            sheet.cell(index, start + offset, item['types'][field])
    total_row = len(records) + 5
    sheet.cell(total_row, start + 3, 'รวม')
    offsets = {**importer.CORE_COLUMNS, **{key: offset for key, _, offset in importer.TYPE_COLUMNS}}
    for field, offset in offsets.items():
        values = [item[field] if field in importer.CORE_COLUMNS else item['types'][field] for item in records]
        known = [value for value in values if value is not None]
        if known:
            sheet.cell(total_row, start + offset, sum(known))
    if extra_summaries:
        for offset in range(28):
            sheet.cell(total_row + 2, start + offset, sheet.cell(total_row, start + offset).value)
        sheet.cell(total_row + 3, start, 'PRIVATE_NOTE_MUST_NOT_BE_PUBLISHED')
    book.save(path)
    book.close()
    return start, total_row


def run_import(path, **kwargs):
    return importer.import_workbook(path, worksheet=kwargs.get('worksheet', 'เขต2'), year=kwargs.get('year', 2569), imported_at=STAMP)


@pytest.mark.parametrize('leading', [False, True])
def test_import_detects_header_offset_excludes_repeated_summaries_and_drops_private_names(tmp_path, leading):
    path = tmp_path / 'synthetic-2569.xlsx'
    _, total = build_workbook(path, leading=leading)
    result = run_import(path)
    assert result['audit']['source_detail_rows'] == 4
    assert result['audit']['excluded_summary_rows'] == [total, total + 2]
    assert sum(group['school_count'] for group in result['cubes']) == 4
    assert sum(group['tested'] or 0 for group in result['cubes']) == 13
    assert sum(group['infected'] or 0 for group in result['cubes']) == 2
    text = json.dumps(result)
    assert 'PRIVATE_' not in text
    assert 'school_name' not in text and 'tambon' not in text and 'name' not in result['cubes'][0]
    assert result['dataset']['year_basis'] == 'filename'
    assert result == run_import(path)


def test_import_preserves_missing_zero_pending_and_source_conflicts(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    build_workbook(path)
    result = run_import(path)
    zero = next(group for group in result['cubes'] if group['status'] == 'zero_tested')
    pending = next(group for group in result['cubes'] if group['status'] == 'pending')
    assert zero['tested'] == 0 and pending['tested'] is None
    assert zero['parasites']['ascaris']['reported'] is None
    assert zero['parasites']['ascaris']['missing'] == 1 and zero['parasites']['ascaris']['tested_missing'] == 0
    assert zero['parasites']['pinworm'] == {'reported': 1, 'consistent': 0, 'conflicting': 1, 'missing': 0, 'tested_missing': 0}
    assert all(values['reported'] is None and values['missing'] == 0 for values in pending['parasites'].values())
    assert result['audit']['missing_count_cells'] == 2
    assert any(issue.get('parasite') == 'pinworm' and issue['code'] == 'source_count_conflict' for issue in result['issues'])


@pytest.mark.parametrize('mutation', ['total', 'header', 'partial_core', 'duplicate', 'sequence'])
def test_bad_source_is_rejected_without_replacing_published_snapshot(tmp_path, monkeypatch, mutation):
    path = tmp_path / 'synthetic-2569.xlsx'
    start, total_row = build_workbook(path)
    book = load_workbook(path)
    sheet = book['เขต2']
    if mutation == 'total':
        sheet.cell(total_row, start + importer.CORE_COLUMNS['tested'], 999)
    elif mutation == 'header':
        sheet.cell(3, start + importer.TYPE_COLUMNS[0][2], 'ชนิดไม่ตรงแม่แบบ')
    elif mutation == 'partial_core':
        sheet.cell(5, start + importer.CORE_COLUMNS['tested']).value = None
    elif mutation == 'duplicate':
        for offset in (0, 1, 2, 3, 27):
            sheet.cell(6, start + offset, sheet.cell(5, start + offset).value)
    elif mutation == 'sequence':
        sheet.cell(6, start - 1, 1)
    book.save(path)
    book.close()
    output = tmp_path / 'published.json'
    sentinel = '{"published": "must survive invalid import"}\n'
    output.write_text(sentinel)
    monkeypatch.setattr(sys, 'argv', ['import_survey', str(path), '--output', str(output), '--worksheet', 'เขต2', '--year', '2569'])
    with pytest.raises(SystemExit) as error:
        importer.main()
    assert error.value.code == 1
    assert output.read_text() == sentinel
    assert not output.with_suffix('.json.tmp').exists()


def test_same_name_location_with_different_affiliation_remains_distinct(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    build_workbook(path, [record(affiliation='หน่วยงาน A'), record(affiliation='หน่วยงาน B')])
    result = run_import(path)
    assert sum(group['school_count'] for group in result['cubes']) == 2
    assert {group['affiliation'] for group in result['cubes']} == {'หน่วยงาน A', 'หน่วยงาน B'}


@pytest.mark.parametrize('year', [2568, 2026])
def test_year_requires_filename_evidence_and_buddhist_calendar(tmp_path, year):
    path = tmp_path / 'synthetic-2569.xlsx'
    build_workbook(path)
    with pytest.raises(ValueError):
        run_import(path, year=year)


def test_truncated_header_is_a_clean_validation_error(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    build_workbook(path)
    book = load_workbook(path)
    sheet = book['เขต2']
    sheet.delete_rows(3, sheet.max_row)
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        run_import(path)


def test_formula_count_without_cached_value_is_not_interpreted_as_pending(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    start, _ = build_workbook(path)
    book = load_workbook(path)
    sheet = book['เขต2']
    # The pending fixture becomes a numeric formula row. openpyxl deliberately
    # does not calculate caches, and the unchanged global totals still reconcile.
    for offset in importer.CORE_COLUMNS.values():
        sheet.cell(7, start + offset, '=0')
    for _, _, offset in importer.TYPE_COLUMNS:
        sheet.cell(7, start + offset, '=0')
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        run_import(path)


def test_different_worksheet_selection_has_a_different_dataset_version(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    start, total_row = build_workbook(path, [record()], extra_summaries=False)
    book = load_workbook(path)
    another = book.copy_worksheet(book['เขต2'])
    another.title = 'อีกเขต'
    another.cell(5, start + importer.CORE_COLUMNS['enrolled'], 20)
    another.cell(total_row, start + importer.CORE_COLUMNS['enrolled'], 20)
    book.save(path)
    book.close()
    first, second = run_import(path), run_import(path, worksheet='อีกเขต')
    assert first['cubes'] != second['cubes']
    assert first['dataset']['version'] != second['dataset']['version']


def test_corrupt_workbook_cli_fails_cleanly_without_overwrite(tmp_path, monkeypatch):
    path = tmp_path / 'synthetic-2569.xlsx'
    path.write_bytes(b'not a valid Office zip archive')
    output = tmp_path / 'published.json'
    output.write_text('{"published":"keep"}')
    monkeypatch.setattr(sys, 'argv', ['import_survey', str(path), '--output', str(output), '--worksheet', 'เขต2', '--year', '2569'])
    with pytest.raises(SystemExit) as error:
        importer.main()
    assert error.value.code == 1 and output.read_text() == '{"published":"keep"}'


def test_positivity_bands_use_exact_boundaries_and_keep_undefined_rates_outside_bins(tmp_path):
    path = tmp_path / 'synthetic-2569.xlsx'
    values = [0, 1, 4, 5, 9, 10, 19, 20, 100]
    records = [record(f'PRIVATE_BOUNDARY_{value}', enrolled=100, tested=100, infected=value) for value in values]
    records += [record('PRIVATE_ZERO', enrolled=0, tested=0, infected=0),
                record('PRIVATE_PENDING', enrolled=None, tested=None, infected=None, coinfection=None)]
    records[-1]['types'] = {key: None for key in records[-1]['types']}
    build_workbook(path, records)
    result = run_import(path)
    counts = {}
    for group in result['cubes']:
        if group['status'] == 'tested':
            counts[group['positivity_band']] = group['school_count']
        else:
            assert group['positivity_band'] is None
    assert counts == {'none': 1, 'under5': 2, '5to10': 2, '10to20': 2, '20plus': 2}
    assert sum(counts.values()) == 9
    assert sum(group['school_count'] for group in result['cubes']) == 11
